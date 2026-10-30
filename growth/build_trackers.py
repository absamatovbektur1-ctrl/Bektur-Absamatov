"""Собирает два трекера роста: с другом и с братишкой.

python3 growth/build_trackers.py  →  growth/*.xlsx
"""
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent
FONT = "Arial"
GREEN, DARK, GOLD, GREY, INPUT = "1F6F4A", "17324D", "C9A227", "F2F2F2", "FFF7D6"
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WD = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]
TOPICS = ["Намаз", "Ният", "Сабр", "Таваккуль", "Тауба", "Ахляк", "Ихсан"]
YES, NO = "✅", "❌"


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def title(ws, text, sub, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws["A1"] = text
    ws["A1"].font = f(True, "FFFFFF", 14)
    ws["A1"].fill = fill(DARK)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A2"] = sub
    ws["A2"].font = f(italic=True, color="555555")


def header(ws, row, cols, color=GREEN, widths=None):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row, i, name)
        c.font = f(True, "FFFFFF")
        c.fill = fill(color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        if widths:
            ws.column_dimensions[L(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 42


def body(ws, r1, r2, c1, c2, wrap=False, center=False, inp=True):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.border = BOX
            if cell.font.name != FONT:
                cell.font = f()
            cell.alignment = Alignment(wrap_text=wrap, vertical="top",
                                       horizontal="center" if center else None)
            if inp:
                cell.fill = fill(INPUT)


def checklist(ws, who, start, days, checks, numbers, notes):
    """Дата | Неделя | День | ✅-колонки | числа | % дня | заметки вечера."""
    cols = ["Дата", "Неделя", "День"] + [c for c, _ in checks] + [n for n, _ in numbers] \
        + ["% дня"] + notes
    width = len(cols)
    title(ws, f"Ежедневный чек-лист — {who}",
          "Ставь ✅ или ❌ в жёлтых клетках (выпадающий список). Числа — вписывай. "
          "% дня считается сам.", width)
    hr = 4
    header(ws, hr, cols, widths=[11, 7, 5] + [9] * len(checks) + [9] * len(numbers)
           + [8] + [30] * len(notes))
    for i, (_, tip) in enumerate(checks + numbers):
        if tip:
            ws.cell(hr, 4 + i).comment = Comment(tip, "Бектур")
    c1, c2 = 4, 3 + len(checks)
    n1 = c2 + 1
    pc = n1 + len(numbers)
    r1, r2 = hr + 1, hr + days
    dv = DataValidation(type="list", formula1=f'"{YES},{NO}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{L(c1)}{r1}:{L(c2)}{r2}")
    num = DataValidation(type="decimal", operator="between", formula1="0", formula2="24",
                         allow_blank=True)
    ws.add_data_validation(num)
    for d in range(days):
        r = r1 + d
        day = start + timedelta(days=d)
        ws.cell(r, 1, day).number_format = "DD.MM"
        ws.cell(r, 2, f"=INT((A{r}-$A${r1})/7)+1")
        ws.cell(r, 3, WD[day.weekday()])
        rng = f"{L(c1)}{r}:{L(c2)}{r}"
        ws.cell(r, pc, f'=IF(COUNTA({rng})=0,"",COUNTIF({rng},"{YES}")/{len(checks)})')
        ws.cell(r, pc).number_format = "0%"
        if day.weekday() == 6:
            for c in range(1, 4):
                ws.cell(r, c).font = f(True, GREEN)
    body(ws, r1, r2, 1, 3, center=True, inp=False)
    body(ws, r1, r2, c1, pc - 1, center=True)
    body(ws, r1, r2, pc, pc, center=True, inp=False)
    body(ws, r1, r2, pc + 1, width, wrap=True)
    for c in range(1, 4):
        for r in range(r1, r2 + 1):
            ws.cell(r, c).fill = fill(GREY)
    cr = f"{L(c1)}{r1}:{L(pc - 1)}{r2}"
    ws.conditional_formatting.add(cr, CellIsRule(operator="equal", formula=[f'"{YES}"'],
                                                 fill=fill("C6EFCE")))
    ws.conditional_formatting.add(cr, CellIsRule(operator="equal", formula=[f'"{NO}"'],
                                                 fill=fill("FFC7CE")))
    ws.conditional_formatting.add(f"{L(pc)}{r1}:{L(pc)}{r2}", ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B",
        mid_type="num", mid_value=0.6, mid_color="FFEB84",
        end_type="num", end_value=1, end_color="63BE7B"))
    for i, (_, minimum) in enumerate(numbers):
        col = L(n1 + i)
        ws.conditional_formatting.add(f"{col}{r1}:{col}{r2}", FormulaRule(
            formula=[f'AND({col}{r1}<>"",{col}{r1}>={minimum})'], fill=fill("C6EFCE")))
        ws.conditional_formatting.add(f"{col}{r1}:{col}{r2}", FormulaRule(
            formula=[f'AND({col}{r1}<>"",{col}{r1}<{minimum})'], fill=fill("FFC7CE")))
    for i in range(len(numbers)):
        num.add(f"{L(n1 + i)}{r1}:{L(n1 + i)}{r2}")
    ws.freeze_panes = ws.cell(r1, 4)
    return dict(r1=r1, r2=r2, c1=c1, c2=c2, n1=n1, pc=pc, checks=checks, numbers=numbers)


def summary(ws, people, weeks, start):
    """people: [(имя, лист, layout)]"""
    cols = ["Неделя", "С", "По"]
    for who, _, lay in people:
        cols += [f"{who}: средний %", f"{who}: намазов из 35"] + \
            [f"{who}: {n.lower()}" for n, _ in lay["numbers"]]
    title(ws, "Итоги по неделям", "Считается автоматически из чек-листов. "
          "Смотрим вместе на созвоне.", len(cols))
    header(ws, 4, cols, color=DARK, widths=[8, 8, 8] + [13] * (len(cols) - 3))
    for w in range(1, weeks + 1):
        r = 4 + w
        ws.cell(r, 1, w)
        ws.cell(r, 2, start + timedelta(days=7 * (w - 1))).number_format = "DD.MM"
        ws.cell(r, 3, start + timedelta(days=7 * w - 1)).number_format = "DD.MM"
        c = 4
        for _, sheet, lay in people:
            s = f"'{sheet}'!"
            wk = f"{s}$B${lay['r1']}:$B${lay['r2']}"
            pc = L(lay["pc"])
            ws.cell(r, c, f'=IFERROR(AVERAGEIFS({s}${pc}${lay["r1"]}:${pc}${lay["r2"]},'
                          f'{wk},$A{r}),"")').number_format = "0%"
            p1, p2 = L(lay["c1"]), L(lay["c1"] + 4)
            ws.cell(r, c + 1, f'=SUMPRODUCT(({wk}=$A{r})*({s}${p1}${lay["r1"]}:${p2}${lay["r2"]}'
                              f'="{YES}"))')
            for i in range(len(lay["numbers"])):
                col = L(lay["n1"] + i)
                ws.cell(r, c + 2 + i, f'=SUMIFS({s}${col}${lay["r1"]}:${col}${lay["r2"]},'
                                      f'{wk},$A{r})')
            c += 2 + len(lay["numbers"])
    body(ws, 5, 4 + weeks, 1, len(cols), center=True, inp=False)
    ws.conditional_formatting.add(f"D5:{L(len(cols))}{4 + weeks}", ColorScaleRule(
        start_type="min", start_color="FFFFFF", end_type="max", end_color="A9D8B8"))
    ws.freeze_panes = "D5"


def calls(ws, names, start, weeks):
    cols = ["№", "Дата (Вс)", "Тема по исламу"] + [f"Отчёт: {n}" for n in names] + \
        ["Ошибки / где были слабы", "Вывод по теме: как применить", "План на неделю", "Созвон был?"]
    title(ws, "Еженедельный созвон — воскресенье, 45–60 мин",
          "1) Отчёт каждого (учёба, намаз, спорт, развитие)  2) Ошибки  "
          "3) Тема по исламу  4) План на неделю", len(cols))
    header(ws, 4, cols, color=DARK, widths=[5, 11, 14] + [32] * len(names) + [30, 32, 32, 10])
    dv = DataValidation(type="list", formula1=f'"{YES},{NO}"', allow_blank=True)
    ws.add_data_validation(dv)
    for w in range(weeks):
        r = 5 + w
        ws.cell(r, 1, w + 1)
        ws.cell(r, 2, start + timedelta(days=7 * w)).number_format = "DD.MM.YYYY"
        ws.cell(r, 3, TOPICS[w % len(TOPICS)])
        ws.row_dimensions[r].height = 60
    n = len(cols)
    body(ws, 5, 4 + weeks, 1, 3, center=True, inp=False)
    body(ws, 5, 4 + weeks, 4, n, wrap=True)
    dv.add(f"{L(n)}5:{L(n)}{4 + weeks}")
    ws.freeze_panes = "D5"


def goals(ws, names):
    title(ws, "Цели", "Месяц — 3 конкретные цели. Год — 3–5 больших. "
          "Цель = что + сколько + к какому числу.", 6)
    header(ws, 4, ["Кто", "Срок", "Сфера", "Цель (конкретно, измеримо)", "Первый шаг",
                   "Статус"], widths=[12, 10, 14, 48, 32, 14])
    dv = DataValidation(type="list", formula1='"Не начато,В процессе,Готово ✅,Провал"',
                        allow_blank=True)
    sphere = DataValidation(type="list",
                            formula1='"Религия,Учёба,Тело,Развитие,Финансы,Характер,Проект"',
                            allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(sphere)
    r = 5
    for who in names:
        for term, k in (("Месяц", 3), ("Год", 5)):
            for _ in range(k):
                ws.cell(r, 1, who)
                ws.cell(r, 2, term)
                r += 1
    body(ws, 5, r - 1, 1, 2, inp=False)
    body(ws, 5, r - 1, 3, 6, wrap=True)
    dv.add(f"F5:F{r - 1}")
    sphere.add(f"C5:C{r - 1}")
    ws["D5"] = "Прочитать 600 страниц за ноябрь (20 стр/день)"
    ws["E5"] = "Сегодня 20 стр. «Сөзлер» после Иша"
    ws["C5"] = "Развитие"
    ws["F5"] = "В процессе"
    for c in "CDEF":
        ws[f"{c}5"].font = f(italic=True, color="7F7F7F")
    ws["D5"].comment = Comment("Пример — замени на свою цель.", "Бектур")


def materials(ws, rows, names):
    title(ws, "Книги, подкасты, навыки", "Отмечай, кто что начал / закончил. "
          "Новые идеи добавляем сюда.", 6 + len(names))
    cols = ["Категория", "Название", "Автор / где", "Зачем", "Формат"] + names + ["Заметки"]
    header(ws, 4, cols, color=DARK, widths=[14, 34, 22, 40, 10] + [12] * len(names) + [30])
    dv = DataValidation(type="list", formula1='"Хочу,Читаю,Прочитал ✅"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            ws.cell(5 + i, 1 + j, v)
    last = 4 + len(rows) + 10
    body(ws, 5, last, 1, 5, wrap=True, inp=False)
    body(ws, 5, last, 6, len(cols), wrap=True)
    dv.add(f"F5:{L(5 + len(names))}{last}")
    ws.freeze_panes = "C5"


def guide(ws, lines):
    ws.column_dimensions["A"].width = 110
    ws["A1"] = lines[0]
    ws["A1"].font = f(True, "FFFFFF", 16)
    ws["A1"].fill = fill(DARK)
    ws.row_dimensions[1].height = 32
    for i, line in enumerate(lines[1:], 3):
        c = ws.cell(i, 1, line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = f(True, GREEN, 12) if line.isupper() or line.startswith("⚡") else f(size=11)
    ws.cell(len(lines) + 3, 1, "Жёлтые клетки — вписываешь ты. Серые и белые — считаются сами, "
            "не трогай.").fill = fill(INPUT)
    ws.sheet_view.showGridLines = False


BOOKS_COMMON = [
    ("Религия", "Рисале-и Нур: «Сөзлер» (Слова)", "Саид Нурси", "Вера через размышление; "
     "начать с «Малых слов» и 23-го Слова", "Книга"),
    ("Религия", "Рисале-и Нур: «Лемалар» (Вспышки)", "Саид Нурси", "Сабр, болезнь, тауба — "
     "прямо к темам созвонов", "Книга"),
    ("Религия", "«Изумрудные холмы сердца»", "Фетхуллах Гюлен", "Ихсан, ахляк, "
     "духовная работа над собой", "Книга"),
    ("Религия", "«Вечный свет: Пророк Мухаммад ﷺ»", "Фетхуллах Гюлен", "Сира — "
     "пример характера и лидерства", "Книга"),
    ("Религия", "Тафсир по 1 странице в день", "любой надёжный тафсир", "Понимать то, "
     "что читаешь в Коране", "Книга"),
    ("Мышление", "«Атомные привычки»", "Джеймс Клир", "Как держать дисциплину — "
     "основа всей этой системы", "Книга"),
    ("Мышление", "«Человек в поисках смысла»", "Виктор Франкл", "Сабр и смысл в трудности",
     "Книга"),
    ("Мышление", "«Думай медленно… решай быстро»", "Даниэль Канеман", "Как мозг ошибается",
     "Книга"),
    ("Учёба", "«Глубокая работа»", "Кэл Ньюпорт", "Фокус 2–4 часа без телефона", "Книга"),
    ("Финансы", "«Самый богатый человек в Вавилоне»", "Джордж Клейсон", "База: копить, "
     "не тратить всё, приумножать", "Книга"),
    ("Финансы", "«Богатый папа, бедный папа»", "Роберт Кийосаки", "Активы vs пассивы", "Книга"),
    ("Характер", "«Как завоёвывать друзей…»", "Дейл Карнеги", "Общение, уважение, "
     "добрые слова", "Книга"),
    ("Подкаст", "Huberman Lab", "YouTube / Spotify", "Сон, спорт, фокус — по науке "
     "(англ., с субтитрами)", "Подкаст"),
    ("Подкаст", "The Diary of a CEO", "YouTube", "Истории людей, бизнес, дисциплина (англ.)",
     "Подкаст"),
    ("Навык", "Английский: 20 мин в день", "Duolingo / Anki / сериалы", "Доступ ко всем "
     "знаниям мира", "Навык"),
    ("Навык", "Публичные выступления", "запись себя на 2 мин/день", "Уверенность, подкаст, "
     "блог", "Навык"),
    ("Навык", "ИИ-инструменты (Claude, ChatGPT)", "практика на учёбе и проектах",
     "Делать за час то, что другие за день", "Навык"),
    ("Навык", "Excel / Google Таблицы", "YouTube-курсы", "Нужен в любой работе", "Навык"),
    ("Игра", "Шахматы: 1 партия/задача в день", "Chess.com / Lichess", "Мышление на шаги "
     "вперёд, можно играть друг с другом онлайн", "Игра"),
]


def build_friend():
    wb = Workbook()
    start = date(2026, 11, 1)
    days, weeks = 63, 9
    names = ["Бектур", "Друг"]
    checks = [("Фаджр", None), ("Зухр", None), ("Аср", None), ("Магриб", None), ("Иша", None),
              ("Коран 10–15 мин", None), ("Зикр", None), ("Без харама", None),
              ("Учёба без телефона", "Фокус-блоки, телефон в другой комнате"),
              ("Развитие 30–60 мин", "Книга / психология / финансы / новости"),
              ("Спорт / активность", None), ("Вода 2 л", None), ("Нормальное питание", None),
              ("Сон 6–8 ч", None), ("Доброе дело", "Помог кому-то"),
              ("Родители", "Позвонил / помог / уважение"),
              ("Без пустой траты времени", "Без бесконечных рилсов и игр"),
              ("Проект / подкаст", "Работа над своим проектом или подкастом")]
    numbers = [("Часов учёбы", 2), ("Страниц книги", 20)]
    notes = ["Что сделал хорошо", "Где был слаб", "Что исправлю завтра"]

    g = wb.active
    g.title = "Как пользоваться"
    guide(g, ["🧱 НАША СИСТЕМА — Бектур и друг", "ОСНОВА",
              "Ислам — фундамент · Дисциплина > настроение · Каждый день рост · Команда = контроль",
              "КАЖДЫЙ ДЕНЬ",
              "1. Открываешь свой лист «Чек-лист — …», ставишь ✅/❌ по каждому пункту.",
              "2. Вписываешь часы учёбы и страницы книги — клетка станет зелёной, если норма "
              "выполнена (2 ч и 20 стр.).",
              "3. Вечером 3 строки самоконтроля: что сделал хорошо, где был слаб, что исправлю.",
              "КАЖДОЕ ВОСКРЕСЕНЬЕ",
              "Созвон 45–60 мин по листу «Созвон»: отчёт → ошибки → тема по исламу → план. "
              "Лист «Итоги» сам покажет средний % за неделю у каждого.",
              "ТЕМЫ ПО КРУГУ",
              "Намаз → Ният → Сабр → Таваккуль → Тауба → Ахляк → Ихсан (уже расставлены по датам).",
              "ПРАВИЛО ЧЕСТНОСТИ",
              "Не ставим ✅ за то, что не сделали. Пропустил — ❌ и пишешь почему. "
              "Если кто-то расслабился — возвращаем, без обид.",
              "ЗАПРЕТЫ",
              "Пропуск намаза · Лень · Пустая трата времени · Слабость характера · "
              "Окружение, которое тянет вниз",
              "⚡ Мы не просто живём — мы себя строим."])

    lays = []
    for n in names:
        ws = wb.create_sheet(f"Чек-лист — {n}")
        lays.append((n, ws.title, checklist(ws, n, start, days, checks, numbers, notes)))
    summary(wb.create_sheet("Итоги"), lays, weeks, start)
    calls(wb.create_sheet("Созвон"), names, start, weeks)
    goals(wb.create_sheet("Цели"), names)
    materials(wb.create_sheet("Книги и навыки"), BOOKS_COMMON, names)

    # пример заполнения первой строки у Бектура
    ws = wb["Чек-лист — Бектур"]
    lay = lays[0][2]
    r = lay["r1"]
    for c in range(lay["c1"], lay["c2"] + 1):
        ws.cell(r, c, YES)
    ws.cell(r, lay["c1"] + 11, NO)
    ws.cell(r, lay["n1"], 3)
    ws.cell(r, lay["n1"] + 1, 25)
    ws.cell(r, lay["pc"] + 1, "Все 5 намазов вовремя, 3 часа учёбы")
    ws.cell(r, lay["pc"] + 2, "Мало воды")
    ws.cell(r, lay["pc"] + 3, "Бутылку воды на стол с утра")
    ws.cell(r, 1).comment = Comment("Строка-пример: перезапиши своими данными.", "Бектур")
    path = OUT / "Трекер_Бектур_и_друг.xlsx"
    wb.save(path)
    return path


def build_brother():
    wb = Workbook()
    start = date(2026, 10, 30)
    days, weeks = 63, 9
    who = "Братишка"
    checks = [("Фаджр", None), ("Зухр", None), ("Аср", None), ("Магриб", None), ("Иша", None),
              ("Коран 10 мин", None), ("Уроки сделаны", "Все домашние задания лицея"),
              ("Задание от Бектура", "Подкаст / видео / книга — с заметками"),
              ("Заметки отправил", "Скрин или фото заметок Бектуру"),
              ("Спорт / активность", None), ("Помог дома", "Родителям, без напоминаний"),
              ("Телефон ≤ 2 ч", "Экранное время в настройках телефона"),
              ("Сон до 23:00", None), ("Доброе дело", None)]
    numbers = [("Часов учёбы", 2), ("Страниц книги", 15)]
    notes = ["Что узнал нового сегодня", "Где был слаб / что исправлю"]

    g = wb.active
    g.title = "Как пользоваться"
    guide(g, ["🧱 ТРЕКЕР РОСТА — братишка (наставник: Бектур)", "КАЖДЫЙ ДЕНЬ",
              "1. Лист «Чек-лист»: ставишь ✅/❌ по каждому пункту честно.",
              "2. Вписываешь часы учёбы и страницы книги (норма: 2 ч и 15 стр.).",
              "3. Пишешь 1 вещь, которую узнал сегодня, и где был слаб.",
              "ЗАДАНИЯ",
              "Лист «Задания»: что задал Бектур. Сделал → заполни заметки по шаблону и отправь "
              "Бектуру скрин. Бектур пишет отзыв и ставит оценку.",
              "ШАБЛОН ЗАМЕТОК ПО ПОДКАСТУ",
              "1) Название и кто говорит  2) 3–5 главных мыслей своими словами  "
              "3) 1 цитата, которая зацепила  4) Что я применю уже завтра  5) С чем не согласен / "
              "вопрос Бектуру",
              "КАЖДОЕ ВОСКРЕСЕНЬЕ",
              "Короткий разговор с Бектуром (15–20 мин): % за неделю на листе «Итоги», "
              "что получилось, тема по исламу, задания на неделю.",
              "⚡ Учёба в лицее — главное. Всё остальное помогает стать сильнее."])

    ws = wb.create_sheet("Чек-лист")
    lay = checklist(ws, who, start, days, checks, numbers, notes)
    summary(wb.create_sheet("Итоги"), [(who, ws.title, lay)], weeks, start)

    t = wb.create_sheet("Задания")
    cols = ["Дата", "Задание", "Ссылка / книга", "Срок", "Заметки брата (главные мысли)",
            "Что применит", "Сдал?", "Отзыв Бектура", "Оценка 1–5"]
    title(t, "Задания от Бектура", "Бектур заполняет A–D и H–I, братишка — E–G.", len(cols))
    header(t, 4, cols, color=DARK, widths=[11, 32, 26, 11, 44, 30, 8, 34, 9])
    plan = [
        (date(2026, 10, 30), "Посмотреть подкаст и выписать заметки по шаблону",
         "ссылку скинет Бектур", date(2026, 10, 30)),
        (date(2026, 11, 1), "«Атомные привычки» — гл. 1–2, 3 главные мысли", "Джеймс Клир",
         date(2026, 11, 3)),
        (date(2026, 11, 3), "Выучить суру и её смысл (Аль-Аср), рассказать своими словами",
         "Коран + тафсир", date(2026, 11, 5)),
        (date(2026, 11, 5), "Подкаст про сон и фокус — что поменяю в режиме",
         "Huberman Lab (с субтитрами)", date(2026, 11, 7)),
        (date(2026, 11, 8), "Составить свои 3 цели на ноябрь (лист «Цели»)", "—",
         date(2026, 11, 8)),
        (date(2026, 11, 10), "Рисале-и Нур: одно из «Малых слов» — пересказ 5 предложений",
         "Саид Нурси, «Сөзлер»", date(2026, 11, 12)),
        (date(2026, 11, 13), "Записать себя на видео 1 мин: что узнал за неделю",
         "камера телефона", date(2026, 11, 14)),
    ]
    for i, row in enumerate(plan):
        for j, v in enumerate(row):
            c = t.cell(5 + i, 1 + j, v)
            if isinstance(v, date):
                c.number_format = "DD.MM"
    last = 4 + len(plan) + 25
    body(t, 5, last, 1, 9, wrap=True)
    for r in range(5, last + 1):
        t.row_dimensions[r].height = 45
    dv = DataValidation(type="list", formula1=f'"{YES},{NO}"', allow_blank=True)
    sc = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        allow_blank=True)
    t.add_data_validation(dv)
    t.add_data_validation(sc)
    dv.add(f"G5:G{last}")
    sc.add(f"I5:I{last}")
    t.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="equal",
                                 formula=[f'"{YES}"'], fill=fill("C6EFCE")))
    t["E5"].comment = Comment("Здесь братишка пишет 3–5 мыслей своими словами.", "Бектур")
    t.freeze_panes = "C5"
    n = 5 + len(plan) + 25 + 2
    t.cell(n, 1, "Сдано заданий:").font = f(True)
    t.cell(n, 2, f'=COUNTIF(G5:G{last},"{YES}")&" из "&COUNTA(B5:B{last})')
    t.cell(n + 1, 1, "Средняя оценка:").font = f(True)
    t.cell(n + 1, 2, f'=IFERROR(ROUND(AVERAGE(I5:I{last}),1),"—")')

    goals(wb.create_sheet("Цели"), [who])
    young = [b for b in BOOKS_COMMON if b[1] not in (
        "«Думай медленно… решай быстро»", "«Богатый папа, бедный папа»")]
    materials(wb.create_sheet("Книги и навыки"), young, [who])
    path = OUT / "Трекер_братишка.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    for p in (build_friend(), build_brother()):
        print(p)
